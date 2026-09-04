import os
import math
from typing import Dict, Any, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from langdetect import detect
except ImportError:
    detect = None


class ExtractedDocument:
    def __init__(
        self,
        text: str,
        title: str,
        page_count: int,
        word_count: int,
        reading_time: int,
        language: str,
        metadata: Dict[str, Any],
        pages: Optional[List[Dict[str, Any]]] = None
    ):
        self.text = text
        self.title = title
        self.page_count = page_count
        self.word_count = word_count
        self.reading_time = reading_time
        self.language = language
        self.metadata = metadata
        self.pages = pages if pages else [{"page_number": 1, "text": text}]


def detect_language(text: str) -> str:
    """Detects text language using langdetect with fallback to 'en'."""
    if not text or len(text.strip()) < 10 or detect is None:
        return "en"
    try:
        lang = detect(text[:2000])
        return str(lang)
    except Exception:
        return "en"


def compute_reading_time(word_count: int) -> int:
    """Calculates estimated reading time in minutes (assumes 200 words per minute)."""
    if word_count <= 0:
        return 1
    return max(1, math.ceil(word_count / 200))


def extract_pdf(file_path: str, default_title: str) -> ExtractedDocument:
    """Extracts text, metadata, page count, and page-by-page chunks from PDF using PyMuPDF (fitz)."""
    if fitz is None:
        raise RuntimeError("PyMuPDF (fitz) library is not installed.")
    
    text_chunks = []
    pages_data = []
    page_count = 0
    doc_title = default_title
    
    doc = fitz.open(file_path)
    page_count = len(doc)
    
    # Try reading PDF metadata title
    pdf_meta = doc.metadata or {}
    if pdf_meta.get("title") and str(pdf_meta["title"]).strip():
        doc_title = str(pdf_meta["title"]).strip()
        
    for page_num, page in enumerate(doc, 1):
        page_text = page.get_text()
        if page_text and page_text.strip():
            clean_page_text = page_text.strip()
            text_chunks.append(clean_page_text)
            pages_data.append({
                "page_number": page_num,
                "text": clean_page_text
            })
            
    doc.close()
    
    full_text = "\n\n".join(text_chunks).strip()
    words = full_text.split()
    word_count = len(words)
    reading_time = compute_reading_time(word_count)
    language = detect_language(full_text)
    
    if not pages_data and full_text:
        pages_data = [{"page_number": 1, "text": full_text}]
        
    return ExtractedDocument(
        text=full_text,
        title=doc_title,
        page_count=page_count,
        word_count=word_count,
        reading_time=reading_time,
        language=language,
        metadata={"format": "PDF", "pdf_metadata": pdf_meta},
        pages=pages_data
    )


def extract_docx(file_path: str, default_title: str) -> ExtractedDocument:
    """Extracts text, metadata, and estimated page groupings from DOCX using python-docx."""
    if docx is None:
        raise RuntimeError("python-docx library is not installed.")
    
    doc = docx.Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs).strip()
    
    words = full_text.split()
    word_count = len(words)
    page_count = max(1, math.ceil(word_count / 350))
    reading_time = compute_reading_time(word_count)
    language = detect_language(full_text)
    
    pages_data = []
    curr_page = 1
    curr_paras = []
    curr_words = 0
    for p in paragraphs:
        curr_paras.append(p)
        curr_words += len(p.split())
        if curr_words >= 350:
            pages_data.append({"page_number": curr_page, "text": "\n\n".join(curr_paras)})
            curr_page += 1
            curr_paras = []
            curr_words = 0
    if curr_paras:
        pages_data.append({"page_number": curr_page, "text": "\n\n".join(curr_paras)})
    
    if not pages_data and full_text:
        pages_data = [{"page_number": 1, "text": full_text}]
        
    return ExtractedDocument(
        text=full_text,
        title=default_title,
        page_count=page_count,
        word_count=word_count,
        reading_time=reading_time,
        language=language,
        metadata={"format": "DOCX", "paragraph_count": len(doc.paragraphs)},
        pages=pages_data
    )


def extract_text_plain(file_path: str, default_title: str, is_markdown: bool = False) -> ExtractedDocument:
    """Reads plain text or markdown file content with encoding fallbacks and section detection."""
    full_text = ""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            full_text = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", errors="replace") as f:
            full_text = f.read()
            
    full_text = full_text.strip()
    words = full_text.split()
    word_count = len(words)
    page_count = max(1, math.ceil(word_count / 400))
    reading_time = compute_reading_time(word_count)
    language = detect_language(full_text)
    
    # Try extracting markdown # title if present
    extracted_title = default_title
    if is_markdown:
        for line in full_text.splitlines():
            line_str = line.strip()
            if line_str.startswith("# "):
                extracted_title = line_str[2:].strip()
                break
                
    # Group text into virtual pages (~400 words per page)
    pages_data = []
    paragraphs = [p.strip() for p in full_text.split("\n\n") if p.strip()]
    curr_page = 1
    curr_paras = []
    curr_words = 0
    for p in paragraphs:
        curr_paras.append(p)
        curr_words += len(p.split())
        if curr_words >= 400:
            pages_data.append({"page_number": curr_page, "text": "\n\n".join(curr_paras)})
            curr_page += 1
            curr_paras = []
            curr_words = 0
    if curr_paras:
        pages_data.append({"page_number": curr_page, "text": "\n\n".join(curr_paras)})
        
    if not pages_data and full_text:
        pages_data = [{"page_number": 1, "text": full_text}]
        
    return ExtractedDocument(
        text=full_text,
        title=extracted_title,
        page_count=page_count,
        word_count=word_count,
        reading_time=reading_time,
        language=language,
        metadata={"format": "Markdown" if is_markdown else "TXT"},
        pages=pages_data
    )


def extract_document(file_path: str, file_extension: str, original_filename: str) -> ExtractedDocument:
    """
    Main extraction dispatcher based on file extension.
    """
    base_name, _ = os.path.splitext(original_filename)
    clean_title = base_name.replace("_", " ").replace("-", " ").title()
    
    ext = file_extension.lower()
    if ext == ".pdf":
        return extract_pdf(file_path, clean_title)
    elif ext == ".docx":
        return extract_docx(file_path, clean_title)
    elif ext == ".txt":
        return extract_text_plain(file_path, clean_title, is_markdown=False)
    elif ext == ".md":
        return extract_text_plain(file_path, clean_title, is_markdown=True)
    else:
        raise ValueError(f"Extraction not supported for file extension '{ext}'")
